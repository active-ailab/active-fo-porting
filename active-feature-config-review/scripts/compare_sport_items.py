#!/usr/bin/env python3
"""Extract owned sport-type/data support, preserving source conflicts and mapping gaps."""
import argparse
import csv
from pathlib import Path
import re

from audit_common import (DEFAULT_CATALOG, bool_entries, catalog, csv_output,
                          fingerprints, json_output)


def one(items, description):
    if len(items) != 1:
        raise ValueError(f'Expected one {description}; found {len(items)}')
    return items[0]


def column(sheet, row, value):
    return one([c.column for c in sheet[row] if c.value == value], f'{sheet.title} row {row}: {value}')


def normalize(value):
    return str(value).strip().replace(' ', '_').replace('-', '_').upper()


def active_sheets(book):
    control = book['code_attr_product']
    first = one([c.row for c in control['A'] if c.value == 'code_attr_product_start'], 'product start')
    last = one([c.row for c in control['A'] if c.value == 'code_attr_product_end'], 'product end')
    return [control.cell(r, 1).value for r in range(first + 1, last) if control.cell(r, 1).value]


def load_project(book, common, raw_csv, product, active):
    sheet_name = one([name for name in active if product.lower() in
        [str(c.value).lower() for c in book[name][2] if c.value]], f'active attribute sheet for {product}')
    sheet = book[sheet_name]
    product_col = one([c.column for c in sheet[2] if str(c.value).lower() == product.lower()], 'product column')
    start = column(sheet, 2, 'product_end') + 1
    end = column(sheet, 2, 'attr_end')
    stop = one([c.row for c in sheet['C'] if c.value == 'SPORT_ATTR_MAX'], 'SPORT_ATTR_MAX')
    attr_rows = list(range(3, stop))
    fields = [sheet.cell(r, 3).value for r in attr_rows]
    ids = [sheet.cell(r, 1).value for r in attr_rows]
    if len(set(fields)) != len(fields) or len(set(ids)) != len(ids) or any(x is None for x in fields + ids):
        raise ValueError('Missing/duplicate attribute identifiers in ' + sheet_name)
    path = common / f'sport_rt_attr_config/sport_rt_attr_{product}.h'
    entries = bool_entries(path)
    if len(entries) != len(attr_rows) * (end - start):
        raise ValueError('Generated header dimension mismatch: ' + str(path))
    exported = [r for r in raw_csv if r.get('Source_Sheet') == sheet_name]
    export_header = one([r for r in exported if 'attr id' in r.values()], 'CSV attribute metadata header') if exported else None
    export_id = one([k for k, v in export_header.items() if v == 'attr id'], 'CSV id column') if export_header else None
    export_product = one([k for k, v in export_header.items() if v.lower() == product.lower()], 'CSV product column') if export_header else None
    export_attrs = {r[export_id]: r for r in exported if export_id and r[export_id].isdigit()}
    matrix = {}
    for gi, col in enumerate(range(start, end)):
        group = sheet.cell(2, col).value
        display = str(sheet.cell(1, col).value)
        group_id = 'SPORT_GROUP_' + normalize(group)
        export_group = one([k for k, v in export_header.items() if v == str(group)], 'CSV group column') if export_header else None
        for ai, row in enumerate(attr_rows):
            enabled, line, comment = entries[gi * len(attr_rows) + ai]
            name, field, attr_id = sheet.cell(row, 2).value, fields[ai], ids[ai]
            if comment != display + ' ' + str(name):
                raise ValueError(f'Generated header order/comment mismatch: {path}:{line}')
            expected = bool(sheet.cell(row, product_col).value and sheet.cell(row, col).value == 'Yes')
            cr = export_attrs.get(str(attr_id))
            csv_enabled = bool(cr.get(export_product) and cr.get(export_group) == 'Yes') if cr is not None else None
            key = (group_id, field)
            if key in matrix:
                raise ValueError('Duplicate normalized group/field: ' + str(key))
            matrix[key] = dict(group=group_id, group_name=display, field=field, attr_id=attr_id, name=name,
                header=enabled, xlsx=expected, csv=csv_enabled, header_line=line, sheet=sheet_name,
                product_cell=sheet.cell(row, product_col).coordinate, group_cell=sheet.cell(row, col).coordinate,
                source_consistent=(enabled == expected == csv_enabled) if csv_enabled is not None else 'CSV_NOT_FOUND')
    return matrix, path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, required=True)
    parser.add_argument('--projects', nargs='+', required=True)
    parser.add_argument('--catalog', type=Path, default=DEFAULT_CATALOG)
    parser.add_argument('--features', nargs='+')
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    if len(set(args.projects)) != len(args.projects):
        parser.error('Projects must be unique')
    from openpyxl import load_workbook
    features = catalog(args.catalog, args.features)
    common = args.root / 'framework/engine/sportEngine/common'
    book_path, csv_path = common / 'sports.xlsx', common / 'sports.csv'
    book = load_workbook(book_path, data_only=True)
    if csv_path.exists():
        with csv_path.open(encoding='utf-8-sig', newline='') as f:
            raw_csv = list(csv.DictReader(f))
    else:
        raw_csv = []
    type_path, group_path = common / 'sport_type_auto.h', common / 'sport_group_auto.c'
    enum = {name: int(value) for name, value in re.findall(r'(SPORT_TYPE_\w+)\s*=\s*(\d+)', type_path.read_text())}
    if not enum:
        raise ValueError('No explicit sport IDs parsed from ' + str(type_path))
    mappings = dict(re.findall(r'DECLARE_INSTACE\(\s*(SPORT_TYPE_\w+)\s*,\s*(SPORT_GROUP_\w+)\s*\)', group_path.read_text()))
    matrices, paths, gaps = {}, [args.catalog, book_path, type_path, group_path], []
    if csv_path.exists():
        paths.append(csv_path)
    active = active_sheets(book)
    for project in args.projects:
        matrices[project], path = load_project(book, common, raw_csv, project, active)
        paths.append(path)
    rows, requested_keys = [], set()
    for feature_id, feature in features.items():
        whole, filtered = set(), set()
        for selector, destination in [('group_sport_types', whole), ('filtered_group_sport_types', filtered)]:
            for sport in feature.get(selector, []):
                if sport not in mappings:
                    gaps.append(dict(feature=feature_id, kind='sport_group_mapping_missing', identifier=sport))
                else:
                    destination.add(mappings[sport])
        for project, matrix in matrices.items():
            present_fields = {key[1] for key in matrix}
            for field in set(feature.get('fields', [])) - present_fields:
                gaps.append(dict(feature=feature_id, project=project, kind='field_missing', identifier=field))
            present_groups = {key[0] for key in matrix}
            for group in (whole | filtered) - present_groups:
                gaps.append(dict(feature=feature_id, project=project, kind='group_missing', identifier=group))
            for key, cell in matrix.items():
                group, field = key
                selected = field in feature.get('fields', []) or group in whole
                if group in filtered:
                    selected |= field in feature.get('group_fields', []) or any(field.startswith(p) for p in feature.get('group_field_prefixes', []))
                if selected:
                    requested_keys.add(key)
                    rows.append(dict(feature=feature_id, project=project, **cell))
    comparisons = []
    for key in sorted(requested_keys):
        values = {project: matrix[key]['header'] if key in matrix else 'MISSING' for project, matrix in matrices.items()}
        comparisons.append(dict(group=key[0], field=key[1], values=values,
            same=(len(set(values.values())) == 1 and 'MISSING' not in values.values())))
    sport_rows = []
    sports = book['Sports']
    id_col = column(sports, 2, 'fw sport id')
    name_col = column(sports, 2, 'cn name of type')
    key_col = column(sports, 1, 'App三方同步筛选Key')
    for project in args.projects:
        project_col = one([c.column for c in sports[1] if str(c.value).lower() == project.lower()], 'Sports project column')
        path = args.root / f'packages/services/sport/include/public/sport_type_in_{project}.h'
        support = bool_entries(path)
        paths.append(path)
        sport_source = {int(sports.cell(r, id_col).value): r for r in range(3, sports.max_row + 1)
                        if isinstance(sports.cell(r, id_col).value, (int, float))}
        for feature_id, feature in features.items():
            for sport in feature.get('sport_types', []):
                sid = enum.get(sport)
                if sid is None or sid not in sport_source or sid >= len(support):
                    gaps.append(dict(feature=feature_id, project=project, kind='sport_type_missing', identifier=sport))
                    continue
                row = sport_source[sid]
                expected = bool(sports.cell(row, key_col).value and sports.cell(row, project_col).value)
                enabled, line, comment = support[sid]
                if comment != str(sports.cell(row, name_col).value):
                    raise ValueError(f'Sport type index/comment mismatch: {path}:{line}')
                sport_rows.append(dict(feature=feature_id, project=project, sport_type=sport, id=sid,
                    group=mappings.get(sport, 'MISSING'), xlsx=expected, header=enabled,
                    source_consistent=enabled == expected, product_cell=sports.cell(row, project_col).coordinate, header_line=line))
    args.out.mkdir(parents=True, exist_ok=True)
    csv_output(args.out / 'data-items.csv', rows, ['feature','project','group','group_name','field','attr_id','name',
        'header','xlsx','csv','header_line','sheet','product_cell','group_cell','source_consistent'])
    csv_output(args.out / 'sport-types.csv', sport_rows, ['feature','project','sport_type','id','group','xlsx','header',
        'source_consistent','product_cell','header_line'])
    conflicts = [r for r in rows if r['source_consistent'] is False]
    incomplete = [r for r in rows if r['source_consistent'] == 'CSV_NOT_FOUND']
    summary = dict(projects=args.projects, features=list(features), unique_group_field_cells=len(requested_keys),
        feature_project_rows=len(rows), comparisons=comparisons,
        differences=[r for r in comparisons if not r['same']], source_conflicts=conflicts,
        incomplete_source_rows=len(incomplete), mapping_gaps=gaps,
        sport_type_conflicts=[r for r in sport_rows if not r['source_consistent']],
        note='Header values describe current attribute support only; type/list reachability, executing core, UI and records require code review.',
        files=fingerprints(paths))
    json_output(args.out / 'data-items.json', summary)
    print(f'Unique group/field cells: {len(requested_keys)}; differences: {len(summary["differences"])}; source conflicts: {len(conflicts)}; mapping gaps: {len(gaps)}')


if __name__ == '__main__':
    main()
